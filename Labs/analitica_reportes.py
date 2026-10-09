"""Indicadores descriptivos y documentación técnica compartida de reportes."""
import io
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from constants import LABORATORIOS, LABS_NAMES


def analizar_registros(tabla, fecha, estado, desde, hasta, categoria=None, asistencia=False):
    """Promedio diario con días sin registros incluidos en el denominador."""
    dias = (hasta - desde).days + 1
    fechas = pd.to_datetime(tabla[fecha], format="mixed", dayfirst=True, errors="coerce")
    iso = tabla[fecha].astype(str).str.match(r"^\d{4}-\d{2}-\d{2}")
    fechas.loc[iso] = pd.to_datetime(tabla.loc[iso, fecha], format="ISO8601", errors="coerce")
    fechas = fechas.dt.normalize()
    indice = pd.date_range(desde, hasta, freq="D")
    diarios = fechas.value_counts().reindex(indice, fill_value=0).sort_index()
    por_dia = pd.DataFrame({"Fecha": diarios.index.strftime("%Y-%m-%d"), "Registros": diarios.values})
    estados = tabla[estado].fillna("Pendiente").replace("", "Pendiente").value_counts()
    series = {"Registros por fecha": por_dia,
              "Distribución por estado": estados.rename_axis("Estado").reset_index(name="Registros")}
    if categoria:
        series["Registros por categoría"] = tabla[categoria].fillna("Sin dato").value_counts().rename_axis(categoria).reset_index(name="Registros")
    indicadores = {"Registros incluidos": len(tabla), "Promedio diario de registros": len(tabla) / dias,
                   "Días del periodo": dias}
    ficha = [
        ("Unidad de análisis", "Un registro del detalle exportado. No equivale necesariamente a una persona única ni a un salón ocupado."),
        ("Periodo", f"{desde.isoformat()} a {hasta.isoformat()}, ambos extremos incluidos"),
        ("Campo temporal", fecha),
        ("Fechas inválidas", "Se excluyen de la serie diaria si no pueden interpretarse. Los filtros de fechas del reporte se aplican antes del cálculo."),
        ("Promedio diario", "Número de registros filtrados / días calendario del periodo. Incluye días sin registros."),
        ("Gráficas", "Recuentos de registros por fecha, estado y categoría, después de aplicar los filtros del reporte."),
        ("Interpretación", "Este reporte describe registros. Para índices de uso de capacidad física consultar Ocupación de salones."),
    ]
    if asistencia:
        valores = tabla[estado].fillna("")
        si = valores.isin(["Si", "Asistio", "Asistió"])
        no = valores.isin(["No", "No asistio", "No asistió"])
        confirmados = int(si.sum() + no.sum())
        indicadores["Asistencia confirmada (%)"] = 100 * int(si.sum()) / confirmados if confirmados else None
        indicadores["Registros pendientes"] = int((~(si | no)).sum())
        ficha.append(("Tasa de asistencia", "100 × registros con asistencia / registros con asistencia o inasistencia confirmada. Pendientes excluidos. Sin confirmaciones: no disponible."))
    return indicadores, series, ficha


def mostrar_analitica(indicadores, series):
    st.markdown("#### Indicadores del periodo")
    cols = st.columns(min(len(indicadores), 4))
    for i, (nombre, valor) in enumerate(indicadores.items()):
        mostrado = "No disponible" if valor is None else (f"{valor:.2f}%" if "(%)" in nombre else f"{valor:.2f}" if isinstance(valor, float) else str(valor))
        cols[i % len(cols)].metric(nombre, mostrado)
    st.caption("Promedio diario calculado sobre todos los días calendario del periodo, incluidos los días sin registros. Las gráficas describen los registros filtrados.")
    for nombre, datos in series.items():
        st.markdown(f"##### {nombre}")
        if len(datos):
            if nombre == "Registros por fecha":
                st.line_chart(datos, x=datos.columns[0], y=datos.columns[1], width="stretch")
            else:
                st.bar_chart(datos, x=datos.columns[0], y=datos.columns[1], width="stretch")


def completar_excel(excel, indicadores=None, series=None, ficha=None, filtros=None):
    """Adjunta especificaciones, capacidades por salón, indicadores y gráficas."""
    libro = load_workbook(io.BytesIO(excel))
    tecnica = libro.create_sheet("Ficha técnica")
    tecnica.append(["Concepto", "Especificación"])
    for nombre, valor in [*(filtros or []), *(ficha or [])]:
        tecnica.append([nombre, str(valor)])
    tecnica.append(["Capacidad física", "Número fijo de bancos de cada salón. No se promedian capacidades entre salones. Ver hoja Capacidades por salón."])
    tecnica.column_dimensions["A"].width = 36
    tecnica.column_dimensions["B"].width = 110
    for row in tecnica.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        tecnica.row_dimensions[row[0].row].height = 45
    capacidades = libro.create_sheet("Capacidades por salón")
    capacidades.append(["Salón", "Bancos físicos", "Unidad"])
    for lab, capacidad in LABORATORIOS.items():
        capacidades.append([LABS_NAMES.get(lab, lab), capacidad, "Bancos simultáneos; capacidad fija del salón"])
    capacidades.column_dimensions["A"].width = 28
    capacidades.column_dimensions["B"].width = 22
    capacidades.column_dimensions["C"].width = 55
    if indicadores is not None:
        hoja = libro.create_sheet("Indicadores")
        hoja.append(["Indicador", "Valor"])
        for nombre, valor in indicadores.items():
            hoja.append([nombre, valor])
        hoja.column_dimensions["A"].width = 45
        hoja.column_dimensions["B"].width = 25
    for i, (nombre, datos) in enumerate((series or {}).items(), 1):
        hoja = libro.create_sheet(f"Análisis {i}")
        for fila in dataframe_to_rows(datos, index=False, header=True):
            hoja.append(fila)
        if len(datos):
            grafica = LineChart() if nombre == "Registros por fecha" else BarChart()
            grafica.title = nombre
            grafica.y_axis.title = "Registros"
            grafica.add_data(Reference(hoja, min_col=2, min_row=1, max_row=hoja.max_row), titles_from_data=True)
            grafica.set_categories(Reference(hoja, min_col=1, min_row=2, max_row=hoja.max_row))
            hoja.add_chart(grafica, "E2")
        hoja.column_dimensions["A"].width = 35
        hoja.column_dimensions["B"].width = 22
    for hoja in libro.worksheets[1:]:
        hoja.freeze_panes = "A2"
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="731116")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
