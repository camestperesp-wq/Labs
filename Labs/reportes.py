# reportes.py

import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
from urllib.parse import urlencode
from constants import LABORATORIOS, HORAS, LABS_NAMES, DIAS
import reservas as res
import prestamos_pasillos as prestamos_pasillos_data
import estudiantes as est
import database as db
import multas
from analitica_reportes import analizar_registros, mostrar_analitica, completar_excel
from ui_components import render_editor_asistencias
from exportaciones import crear_excel_institucional as _crear_excel_institucional


def _render_editor_paginado(df, key, laboratorio=None, tamano=5):
    pagina_key = f"pagina_{key}"
    total = len(df)
    total_paginas = max(1, (total + tamano - 1) // tamano)
    pagina = min(max(1, int(st.session_state.get(pagina_key, 1))), total_paginas)
    st.session_state[pagina_key] = pagina
    inicio = (pagina - 1) * tamano
    render_editor_asistencias(df.iloc[inicio:inicio + tamano].copy(), f"{key}_p{pagina}", laboratorio)

    anterior, info, siguiente = st.columns([1, 2, 1])
    with anterior:
        if st.button("Anterior", key=f"anterior_{key}", disabled=pagina == 1, use_container_width=True):
            st.session_state[pagina_key] = pagina - 1
            st.rerun()
    info.markdown(
        f"<div style='text-align:center;padding:.55rem;color:#5f6368'>Registros {inicio + 1}–{min(inicio + tamano, total)} de {total}</div>",
        unsafe_allow_html=True,
    )
    with siguiente:
        if st.button("Siguiente", key=f"siguiente_{key}", disabled=pagina == total_paginas, use_container_width=True):
            st.session_state[pagina_key] = pagina + 1
            st.rerun()


def mostrar_consulta_fecha_lab():
    st.subheader("Consultar por día, laboratorio y hora")
    fecha = st.date_input("Fecha de la reserva", datetime.now().date(), key="labs_fecha_consulta", help="Consulta únicamente las reservas de este día.")
    lab_cons = st.selectbox("Salón de la reserva", list(LABORATORIOS.keys()), key="labs_lab_consulta", format_func=lambda x: LABS_NAMES.get(x, x), help="Solo incluye las reservas del salón seleccionado.")
    hora_cons = st.selectbox("Bloque horario de la reserva", ["Todas"] + HORAS, key="labs_hora_consulta", help="Todas incluye todos los bloques de la fecha; un bloque limita la consulta a esas dos horas.")

    if st.button("Buscar", key="labs_buscar_fecha_lab"):
        st.session_state.pagina_labs_fecha_lab = 1
        st.session_state.labs_params = {
            "fecha": fecha,
            "lab": lab_cons,
            "hora": hora_cons
        }
        st.rerun()

    if "labs_params" in st.session_state:
        params = st.session_state.labs_params
        fecha_str = params["fecha"].strftime("%Y-%m-%d")
        
        if params["hora"] == "Todas":
            df = res.get_reservas_fecha_lab(fecha_str, params["lab"])
        else:
            df = res.get_reservas_fecha_lab_hora(fecha_str, params["lab"], params["hora"])
            if 'hora' in df.columns:
                df = df.drop(columns=['hora'])

        if df.empty:
            st.info(f"No hay reservas para el {params['fecha'].strftime('%d/%m/%Y')} en {params['lab']}" + 
                    (f" a las {params['hora']}" if params['hora'] != "Todas" else ""))
        else:
            st.success(f"Reservas: {len(df)} encontradas")
            _render_editor_paginado(df, "labs_fecha_lab", params["lab"])

def mostrar_busqueda_codigo():
    st.subheader("Buscar por nombre, código, cédula o QR")
    with st.form("form_buscar_codigo", border=False):
        buscar_col, boton_col = st.columns([5, 1], vertical_alignment="bottom")
        with buscar_col:
            termino = st.text_input("Nombres y apellidos, código, cédula o QR", key="labs_termino_persona")
        with boton_col:
            buscar = st.form_submit_button("Buscar", use_container_width=True)

    if buscar:
        termino = est.normalizar_entrada_busqueda(termino)
        if termino and len(termino) >= 3:
            st.session_state.labs_codigo_busqueda = est.resolver_codigo(termino)
            st.session_state.pagina_labs_persona = 1
            st.rerun()
        else:
            st.warning("Ingresa al menos 3 caracteres")

    if "labs_codigo_busqueda" in st.session_state:
        termino = st.session_state.labs_codigo_busqueda
        coincidencias = est.buscar_personas(termino)
        exactas = coincidencias[(coincidencias.codigo == termino) | (coincidencias.documento == termino)]
        if not exactas.empty:
            termino = str(exactas.iloc[0].codigo)
        elif not coincidencias.empty:
            opciones = dict(zip(coincidencias.codigo, coincidencias.nombres))
            termino = st.selectbox("Usuario encontrado", list(opciones),
                                   format_func=lambda codigo: f"{opciones[codigo]} ({codigo})",
                                   key=f"usuario_reserva_{termino}")
        historial_multas = multas.obtener_multas_estudiante(termino)
        prestamos_activos = prestamos_pasillos_data.obtener_prestamos_activos_codigo(termino)
        df_persona = res.buscar_reservas_persona(termino)
        solicitante = prestamos_pasillos_data.obtener_solicitante(termino)
        nombre = (
            str(df_persona.iloc[0]["nombres"])
            if not df_persona.empty else
            str(prestamos_activos.iloc[0]["solicitante_nombre"])
            if not prestamos_activos.empty else
            str(solicitante["nombres"])
            if solicitante else ""
        )

        if nombre:
            st.success(f"Usuario verificado: {nombre}")
        elif df_persona.empty and prestamos_activos.empty and historial_multas.empty:
            st.info("No se encontraron datos para el código consultado.")
            return

        st.subheader("Informacion del Usuario / Reserva Basica")
        proyecto = ""
        if not df_persona.empty:
            proyecto = str(df_persona.iloc[0].get("proyecto") or "")
        elif solicitante:
            proyecto = str(solicitante.get("proyecto") or "")
        st.dataframe(
            pd.DataFrame([{"Código": termino, "Nombre": nombre or "No registrado", "Proyecto": proyecto}]),
            use_container_width=True,
            hide_index=True,
        )

        st.subheader("Multas activas")
        multas_activas = int(historial_multas["pagado"].eq("NO").sum())
        if multas_activas:
            st.warning(f"El usuario tiene {multas_activas} multa(s) activa(s).")
            destino_deudores = "/deudores?" + urlencode({
                "codigo_deudor": termino,
            })
            st.markdown(
                f'<a href="{destino_deudores}" target="_self">Ver detalle en Deudores</a>',
                unsafe_allow_html=True,
            )
        else:
            st.caption("Sin multas activas registradas.")

        if not historial_multas.empty:
            st.subheader("Historial de multas")
            st.dataframe(historial_multas[["fecha_multa", "motivo", "sancion", "observaciones", "pagado", "fecha_pago", "tecnico_asigna"]].rename(columns={
                "fecha_multa": "Fecha", "motivo": "Motivo", "sancion": "Sanción",
                "observaciones": "Observaciones", "pagado": "Pagado", "fecha_pago": "Fecha de pago",
                "tecnico_asigna": "Técnico",
            }), hide_index=True, width="stretch")

        resumen = pd.DataFrame()
        if not df_persona.empty:
            reservas = df_persona.copy()
            reservas["_fecha"] = pd.to_datetime(reservas["fecha"], errors="coerce")
            reservas["laboratorio"] = reservas["laboratorio"].map(
                lambda salon: LABS_NAMES.get(salon, salon)
            )
            reservas["fecha"] = reservas["_fecha"].dt.strftime("%d/%m/%Y").fillna(reservas["fecha"])
            reservas["multas_activas"] = reservas["multas_activas"].fillna(0).astype(int)
            reservas["Asistencia"] = "Pendiente"
            resumen = reservas.rename(columns={
                "codigo": "Codigo",
                "nombres": "Nombre",
                "proyecto": "Proyecto",
                "fecha": "Fecha",
                "laboratorio": "Salon",
                "hora": "Hora",
            })

            st.subheader("Horario y Detalles de la Clase")
            tabla_asistencia = resumen[["id", "Salon", "Fecha", "Hora", "Asistencia"]].copy()
            tabla_editada = st.data_editor(
                tabla_asistencia,
                key=f"busqueda_asistencia_editor_{termino}",
                hide_index=True,
                use_container_width=True,
                disabled=["id", "Salon", "Fecha", "Hora"],
                column_config={
                    "id": None,
                    "Asistencia": st.column_config.SelectboxColumn(
                        "Asistencia",
                        options=["Pendiente", "Asistió", "No asistió"],
                        required=True,
                    ),
                },
            )
            cambios = tabla_editada.merge(
                tabla_asistencia[["id", "Asistencia"]],
                on="id",
                suffixes=("_nuevo", "_anterior"),
            )
            for _, cambio in cambios.iterrows():
                nuevo_estado = str(cambio["Asistencia_nuevo"] or "Pendiente")
                estado_anterior = str(cambio["Asistencia_anterior"] or "Pendiente")
                if nuevo_estado == estado_anterior or nuevo_estado == "Pendiente":
                    continue
                estado_db = "Si" if nuevo_estado == "Asistió" else "No"
                res.actualizar_asiste(int(cambio["id"]), estado_db, None)
                st.rerun()
        else:
            st.subheader("Horario y Detalles de la Clase")
            st.info("El usuario no tiene una reserva vigente o seleccionable.")

        st.subheader("Prestamos Activos Vigentes")
        if prestamos_activos.empty:
            st.info("El usuario no tiene préstamos activos.")
        else:
            for _, prestamo in prestamos_activos.iterrows():
                with st.container(border=True):
                    detalle_col, accion_col = st.columns([5, 1.6], vertical_alignment="center")
                    detalle_col.markdown(
                        f"**{prestamo['equipos']}**  \nSalida: {prestamo['fecha_salida']}"
                    )
                    destino = "/prestamos?" + urlencode({
                        "prestamo_id": int(prestamo["id"]),
                        "codigo_prestamo": str(prestamo["solicitante"]),
                    })
                    accion_col.markdown(
                        f'<a class="labs-action-link" href="{destino}" target="_self">'
                        "Gestionar devolución</a>",
                        unsafe_allow_html=True,
                    )

def obtener_reporte_ocupacion(fecha_desde, fecha_hasta):
    """Proyecta el horario vigente y los registros del día, una fila por salón/bloque."""
    import calendario as cal
    if fecha_desde > fecha_hasta:
        raise ValueError("La fecha inicial no puede ser posterior a la final.")
    filas = []
    fecha = fecha_desde
    while fecha <= fecha_hasta:
        fecha_str = fecha.isoformat()
        dia = DIAS[fecha.weekday()] if fecha.weekday() < len(DIAS) else "Domingo"
        contexto = cal._build_contexto_calendario(dia, fecha_str)
        mapa = res.obtener_intercambios_fecha(fecha_str)
        for hora in HORAS:
            for lab in LABORATORIOS:
                origen = mapa.get(f"{hora}|{lab}", f"{hora}|{lab}")
                hora_base, lab_base = origen.split("|", 1)
                estado = cal._obtener_estado_celda(dia, lab_base, fecha_str, hora_base, contexto)
                if (hora_base, lab_base) != (hora, lab):
                    estado = cal._ajustar_capacidad_destino(estado, fecha_str, hora_base, lab_base, lab)
                horario = estado.get("horario") or {}
                adicional = cal._es_bloque_reservable(horario.get("asignatura"), horario.get("carrera"))
                if adicional:
                    tipo = "Práctica libre" if "práctica" in str(horario.get("carrera", "")).casefold() or "practica" in str(horario.get("asignatura", "")).casefold() else "Adicional"
                elif estado["tiene_asignatura"]:
                    tipo = "Clase fija"
                elif estado["tiene_profesor"]:
                    tipo = "Reserva docente"
                elif estado["reservas_activas"]:
                    tipo = "Reserva individual"
                else:
                    tipo = "Libre"
                docente = estado["profesor_nombre"] or horario.get("profesor") or ""
                asistencia = {"Si": "Asistió", "No": "No asistió"}.get(
                    estado["estado_profesor"], "Pendiente" if docente else "No aplica"
                )
                filas.append({
                    "Fecha": fecha_str, "Hora": hora, "Salón": LABS_NAMES.get(lab, lab),
                    "Tipo": tipo, "Actividad": horario.get("asignatura") or estado["profesor_asignatura"] or "",
                    "Docente": docente, "Monitor": estado.get("monitor") or "",
                    "Asistencia docente": asistencia, "Capacidad física": estado["total"],
                    "Bancos ocupados o bloqueados": estado["ocupados"], "Bancos disponibles": estado["disponibles"],
                })
        fecha += timedelta(days=1)
    return pd.DataFrame(filas)


def resumir_ocupacion(tabla):
    """Índices ponderados por capacidad, sobre todos los bloques seleccionados."""
    capacidad = tabla["Capacidad física"].sum()
    ocupados = tabla["Bancos ocupados o bloqueados"].sum()
    confirmados = tabla["Asistencia docente"].isin(["Asistió", "No asistió"])
    resumen = {
        "Índice de ocupación (%)": 100 * ocupados / capacidad if capacidad else 0,
        "Bloques con ocupación (%)": 100 * tabla["Bancos ocupados o bloqueados"].gt(0).mean() if len(tabla) else 0,
        "Promedio de bancos por bloque": ocupados / len(tabla) if len(tabla) else 0,
        "Asistencia docente (%)": 100 * tabla["Asistencia docente"].eq("Asistió").sum() / confirmados.sum() if confirmados.any() else None,
        "Asistencias docentes confirmadas": int(confirmados.sum()),
        "Asistencias docentes pendientes": int(tabla["Asistencia docente"].eq("Pendiente").sum()),
        "Capacidad acumulada (banco-bloques de 2 h)": int(capacidad),
    }
    series = {}
    for columna in ("Salón", "Fecha", "Hora"):
        agrupado = tabla.groupby(columna, as_index=False)[["Capacidad física", "Bancos ocupados o bloqueados"]].sum()
        agrupado["Ocupación (%)"] = (100 * agrupado["Bancos ocupados o bloqueados"] / agrupado["Capacidad física"].replace(0, float("nan"))).fillna(0)
        agrupado = agrupado.rename(columns={
            "Capacidad física": "Capacidad acumulada (banco-bloques de 2 h)",
            "Bancos ocupados o bloqueados": "Ocupación acumulada (banco-bloques de 2 h)",
        })
        series[columna] = agrupado
    return resumen, series


def _agregar_estadisticas_excel(excel, resumen, series):
    import io
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils.dataframe import dataframe_to_rows
    libro = load_workbook(io.BytesIO(excel))
    hoja = libro.create_sheet("Indicadores")
    hoja.append(["Indicador", "Valor"])
    for nombre, valor in resumen.items():
        hoja.append([nombre, None if valor is None else float(valor)])
    hoja.append(["Base de cálculo", "Todos los bloques del rango y salón seleccionados, incluidos los libres."])
    hoja.append(["Unidad de capacidad acumulada", "Un banco disponible durante un bloque de 2 horas. No es el número de bancos físicos distintos."])
    hoja.append(["Horario analizado", f"{HORAS[0][:5]} a {HORAS[-1][-5:]} · {len(HORAS)} bloques de 2 horas por día, incluidos todos los días del rango."])
    hoja.append(["Interpretación", "Ocupación programada o bloqueada; no equivale a presencia confirmada."])
    hoja.append(["Asistencia docente", "Asistió / (Asistió + No asistió); pendientes excluidos."])
    hoja.column_dimensions["A"].width = 42
    hoja.column_dimensions["B"].width = 90
    for columna, datos in series.items():
        hoja = libro.create_sheet(f"Por {columna.lower()}")
        for fila in dataframe_to_rows(datos, index=False, header=True):
            hoja.append(fila)
        grafica = LineChart() if columna == "Fecha" else BarChart()
        grafica.title = f"Ocupación por {columna.lower()}"
        grafica.y_axis.title = "Ocupación (%)"
        grafica.x_axis.title = columna
        grafica.add_data(Reference(hoja, min_col=4, min_row=1, max_row=hoja.max_row), titles_from_data=True)
        grafica.set_categories(Reference(hoja, min_col=1, min_row=2, max_row=hoja.max_row))
        hoja.add_chart(grafica, "F2")
        for letra in "ABCD":
            hoja.column_dimensions[letra].width = 28
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def mostrar_reporte_ocupacion():
    st.subheader("Ocupación de salones y asistencia docente")
    st.caption("Incluye clases fijas, adicionales, prácticas libres y reservas. Una fila representa un salón y bloque de dos horas. Los bancos ocupados o bloqueados incluyen la ocupación programada: no equivalen a personas cuya asistencia fue confirmada.")
    with st.form("filtros_ocupacion_salones"):
        desde_col, hasta_col = st.columns(2)
        desde = desde_col.date_input("Fecha inicial (incluida)", datetime.now().date(),
                                     help="Proyecta el horario semanal actual al rango elegido; no reconstruye versiones anteriores del horario.")
        hasta = hasta_col.date_input("Fecha final (incluida)", datetime.now().date())
        salon = st.selectbox("Salón incluido", ["Todos", *LABORATORIOS],
                             format_func=lambda valor: "Todos los salones" if valor == "Todos" else LABS_NAMES.get(valor, valor))
        incluir_libres = st.checkbox("Incluir bloques sin clase, adicional ni reserva", value=False)
        generar = st.form_submit_button("Generar reporte de ocupación")
    if generar:
        if desde > hasta:
            st.error("La fecha inicial no puede ser posterior a la final.")
            return
        st.session_state.reporte_ocupacion_filtros = (desde, hasta, salon, incluir_libres)
    parametros = st.session_state.get("reporte_ocupacion_filtros")
    if not parametros:
        return
    desde, hasta, salon, incluir_libres = parametros
    tabla = obtener_reporte_ocupacion(desde, hasta)
    if salon != "Todos":
        tabla = tabla[tabla["Salón"] == LABS_NAMES.get(salon, salon)]
    resumen, series = resumir_ocupacion(tabla)
    bancos_fisicos = sum(LABORATORIOS.values()) if salon == "Todos" else LABORATORIOS[salon]
    dias = (hasta - desde).days + 1
    capacidad_acumulada = resumen["Capacidad acumulada (banco-bloques de 2 h)"]
    st.caption(f"Base de cálculo: {dias} días calendario · {HORAS[0][:5]}–{HORAS[-1][-5:]} · {len(HORAS)} bloques diarios de 2 h. Capacidad simultánea del conjunto seleccionado: {bancos_fisicos} bancos. Capacidad acumulada: {capacidad_acumulada} banco-bloques de 2 h; ponderada por la capacidad individual de cada salón.")
    with st.expander("Capacidades por salón y metodología"):
        st.table(pd.DataFrame([
            {"Salón": LABS_NAMES.get(lab, lab), "Bancos físicos": capacidad}
            for lab, capacidad in LABORATORIOS.items() if salon == "Todos" or lab == salon
        ]))
        st.markdown("**Ocupación global (%)** = 100 × Σ bancos ocupados o bloqueados por salón y bloque / Σ capacidad física de cada salón y bloque. Es una tasa ponderada por capacidad, no el promedio simple de los porcentajes de los salones.")
    tabla["Ocupación (%)"] = (100 * tabla["Bancos ocupados o bloqueados"] / tabla["Capacidad física"].replace(0, float("nan"))).fillna(0).round(2)
    st.caption("Los indicadores incluyen todos los bloques del rango y salón elegidos, también los libres. La casilla de bloques libres solo cambia las filas visibles del detalle.")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Índice de ocupación", f"{resumen['Índice de ocupación (%)']:.1f}%",
              help="Suma de bancos ocupados o bloqueados / suma de capacidad física de todos los bloques. Ponderado por capacidad.")
    c2.metric("Bloques con ocupación", f"{resumen['Bloques con ocupación (%)']:.1f}%",
              help="Bloques con al menos un banco ocupado o bloqueado / todos los bloques disponibles.")
    c3.metric("Bancos por bloque", f"{resumen['Promedio de bancos por bloque']:.1f}",
              help="Promedio de bancos ocupados o bloqueados por salón y bloque de dos horas, incluidos los libres.")
    asistencia = resumen["Asistencia docente (%)"]
    c4.metric("Asistencia docente", "Sin confirmar" if asistencia is None else f"{asistencia:.1f}%",
              help="Asistió / (Asistió + No asistió). Las asistencias pendientes no entran en este porcentaje.")
    st.caption(f"Asistencia docente: {resumen['Asistencias docentes confirmadas']} registros confirmados y {resumen['Asistencias docentes pendientes']} pendientes.")
    st.markdown("#### Índices de ocupación")
    st.caption("Porcentaje de bancos ocupados o bloqueados respecto de la capacidad disponible en cada agrupación.")
    st.bar_chart(series["Salón"], x="Salón", y="Ocupación (%)", width="stretch")
    st.line_chart(series["Fecha"], x="Fecha", y="Ocupación (%)", width="stretch")
    st.bar_chart(series["Hora"], x="Hora", y="Ocupación (%)", width="stretch")
    if not incluir_libres:
        tabla = tabla[tabla["Tipo"] != "Libre"]
    st.caption(f"Resultado generado: {desde:%d/%m/%Y} al {hasta:%d/%m/%Y} · {LABS_NAMES.get(salon, 'Todos los salones')}")
    if tabla.empty:
        st.info("No hay bloques que coincidan con estos filtros.")
        return
    detalle = tabla.rename(columns={"Capacidad física": "Bancos físicos del salón (por bloque)"})
    st.dataframe(detalle, hide_index=True, width="stretch",
                 column_config={"Bancos físicos del salón (por bloque)": st.column_config.NumberColumn(
                     help="Cantidad fija de bancos del salón de esta fila, durante el bloque indicado en Hora. No es una suma por día.")})
    excel = _crear_excel_institucional(
        detalle, "Ocupación de salones y asistencia docente",
        "Horario vigente y registros de reservas; ocupación programada y asistencia registrada",
        [("Periodo", f"{desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"),
         ("Salón", LABS_NAMES.get(salon, "Todos los salones")), ("Bloques incluidos", len(tabla))],
    )
    excel = _agregar_estadisticas_excel(excel, resumen, series)
    excel = completar_excel(excel, ficha=[
        ("Unidad de análisis", "Un salón en un bloque de dos horas. Capacidad individual según hoja Capacidades por salón."),
        ("Horario analizado", f"{HORAS[0][:5]}–{HORAS[-1][-5:]}; {len(HORAS)} bloques diarios"),
        ("Días analizados", f"{dias} días calendario, incluidos fines de semana y días sin registros"),
        ("Ocupación global (%)", "100 × suma de bancos ocupados o bloqueados en cada salón-bloque / suma de capacidades individuales en cada salón-bloque. Ponderación por capacidad; no promedio simple entre salones."),
        ("Capacidad acumulada", f"{capacidad_acumulada} banco-bloques de 2 h. Cada capacidad individual se multiplica por los bloques y días del periodo."),
        ("Bloques con ocupación (%)", "100 × bloques con al menos un banco ocupado o bloqueado / total de salón-bloques seleccionados."),
        ("Promedio de bancos por bloque", "Suma de bancos ocupados o bloqueados / total de salón-bloques, incluidos los libres."),
        ("Asistencia docente (%)", "100 × Asistió / (Asistió + No asistió). Pendientes excluidos; sin confirmaciones: no disponible."),
        ("Fuente y alcance", "Horario semanal vigente proyectado al periodo y registros activos de reservas. No reconstruye versiones históricas. Incluye intercambios temporales de la sesión actual."),
        ("Detalle y agregaciones", "La opción de ocultar bloques libres afecta solo el detalle exportado. Indicadores y gráficas mantienen todos los bloques del rango y salones seleccionados."),
        ("Ocupación y presencia", "La ocupación programada o bloqueada no equivale a personas con asistencia confirmada."),
    ], filtros=[("Periodo", f"{desde} a {hasta}"), ("Salón", LABS_NAMES.get(salon, "Todos")),
                ("Detalle incluye bloques libres", incluir_libres)])
    st.download_button("Descargar ocupación en Excel", excel,
                       f"ocupacion_salones_{desde:%Y%m%d}_{hasta:%Y%m%d}.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def mostrar_reporte_completo():
    st.subheader("Reporte completo de reservas")
    st.caption("Incluye reservas activas por fecha programada, con ambos extremos del rango incluidos. Asistio: asistencia confirmada; No asistio: ausencia registrada; Pendiente: sin confirmación. No incluye reservas canceladas.")
    
    c1, c2 = st.columns(2)
    with c1:
        fecha_desde = st.date_input("Fecha inicial de reservas (incluida)", datetime.now().date() - timedelta(days=30), key="labs_reporte_desde", help="Filtra por la fecha programada de la reserva, no por la fecha en que se registró.")
    with c2:
        fecha_hasta = st.date_input("Fecha final de reservas (incluida)", datetime.now().date(), key="labs_reporte_hasta", help="Incluye las reservas hasta este día completo.")
    
    # Filtro adicional para tipo de reserva
    tipo_reserva = st.selectbox(
        "Tipo de reserva",
        ["Todas", "Estudiantes (bancos individuales)", "Docentes (sala completa)"],
        key="labs_reporte_tipo",
        help="Todas: bancos individuales y salas completas. Estudiantes: banco mayor que cero. Docentes: reserva de sala completa (banco cero)."
    )
    
    if st.button("Generar reporte", key="labs_generar_reporte"):
        if fecha_desde > fecha_hasta:
            st.error("Fecha 'Desde' > 'Hasta'")
        else:
            df = res.get_reporte_completo(fecha_desde.strftime("%Y-%m-%d"), fecha_hasta.strftime("%Y-%m-%d"))
            
            if df.empty:
                st.info("Sin reservas en el rango.")
                return
            
            # Aplicar filtro por tipo de reserva
            if tipo_reserva == "Estudiantes (bancos individuales)":
                df = df[df['banco'] > 0]
            elif tipo_reserva == "Docentes (sala completa)":
                df = df[df['banco'] == 0]
            # "Todas" no aplica filtro
            
            if df.empty:
                st.info("No hay reservas de este tipo en el rango seleccionado.")
                return
            
            # Mostrar el reporte
            st.dataframe(df, use_container_width=True, hide_index=True)
            etiquetas = {
                "fecha": "Fecha", "hora": "Hora", "laboratorio": "Laboratorio",
                "banco": "Banco", "codigo": "Código", "nombres": "Persona",
                "proyecto": "Programa / Asignatura", "observaciones": "Observaciones",
                "tecnico": "Técnico", "estado": "Estado",
            }
            df_exportar = df.rename(columns=etiquetas)
            indicadores, series, ficha = analizar_registros(df_exportar, "Fecha", "Estado", fecha_desde, fecha_hasta, "Laboratorio", asistencia=True)
            mostrar_analitica(indicadores, series)
            excel = _crear_excel_institucional(
                df_exportar,
                "Reporte ejecutivo de reservas",
                "Control de ocupación y asistencia de laboratorios",
                [
                    ("Periodo", f"{fecha_desde.strftime('%d/%m/%Y')} al {fecha_hasta.strftime('%d/%m/%Y')}"),
                    ("Tipo de reserva", tipo_reserva),
                    ("Total de registros", len(df_exportar)),
                ],
            )
            excel = completar_excel(excel, indicadores, series, ficha,
                                   filtros=[("Tipo de reserva", tipo_reserva), ("Fuente", "Reservas activas; no incluye clases fijas sin registro de asistencia o reserva.")])
            st.download_button(
                "Descargar Excel institucional",
                excel,
                f"reporte_reservas_{fecha_desde.strftime('%Y%m%d')}_{fecha_hasta.strftime('%Y%m%d')}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="labs_descargar_reporte",
            )
            
            # Mostrar resumen estadístico
            st.subheader(" Resumen del reporte")
            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("Total de reservas", len(df))
            with col2:
                asistencias = len(df[df['estado'] == 'Asistio'])
                st.metric("Asistieron", asistencias)
            with col3:
                no_asistencias = len(df[df['estado'] == 'No asistio'])
                st.metric("No asistieron", no_asistencias)

def mostrar_reporte_asistencia_docentes():
    """
    Reporte específico para asistencia de docentes.
    Filtra reservas con banco = 0 (sala completa) y codigo = 'PROFESOR'.
    """
    st.subheader(" Reporte de asistencia de docentes")
    st.caption("Incluye reservas activas de sala completa con código PROFESOR. Si: asistencia confirmada; No: inasistencia registrada; vacío: pendiente de confirmar. El rango incluye ambos días.")
    
    c1, c2 = st.columns(2)
    with c1:
        fecha_desde = st.date_input(
            "Desde",
            datetime.now().date() - timedelta(days=30),
            key="doc_reporte_desde", help="Primera fecha programada de reserva incluida en el informe."
        )
    with c2:
        fecha_hasta = st.date_input(
            "Hasta",
            datetime.now().date(),
            key="doc_reporte_hasta", help="Última fecha programada de reserva incluida en el informe."
        )
    
    # Filtro por laboratorio
    lab_filter = st.selectbox(
        "Laboratorio",
        ["Todos"] + list(LABORATORIOS.keys()),
        format_func=lambda x: "Todos los salones" if x == "Todos" else LABS_NAMES.get(x, x),
        key="doc_reporte_lab", help="Todos los salones incluye todas las reservas docentes; un salón limita el informe a ese espacio."
    )
    
    if st.button("Generar reporte docentes", key="doc_generar_reporte"):
        if fecha_desde > fecha_hasta:
            st.error("Fecha 'Desde' > 'Hasta'")
            return
        
        # Obtener solo registros de docentes (banco = 0, codigo = 'PROFESOR')
        df = res.get_reporte_docentes(
            fecha_desde.strftime("%Y-%m-%d"),
            fecha_hasta.strftime("%Y-%m-%d")
        )
        
        if df.empty:
            st.info("No hay registros de asistencia de docentes en el rango seleccionado.")
            return
        
        # Aplicar filtro de laboratorio
        if lab_filter != "Todos":
            df = df[df['laboratorio'] == lab_filter]
        
        if df.empty:
            st.info(f"No hay registros para el laboratorio seleccionado en este rango.")
            return
        
        # Mostrar el reporte
        st.dataframe(
            df[['fecha', 'hora', 'laboratorio', 'nombres', 'proyecto', 'asiste', 'tecnico', 'observaciones']],
            column_config={
                "fecha": "Fecha",
                "hora": "Hora",
                "laboratorio": "Laboratorio",
                "nombres": "Docente",
                "proyecto": "Asignatura/Motivo",
                "asiste": "Estado",
                "tecnico": "Técnico",
                "observaciones": "Observaciones"
            },
            use_container_width=True,
            hide_index=True
        )
        
        columnas_reporte = {
            "fecha": "Fecha", "hora": "Hora", "laboratorio": "Laboratorio",
            "nombres": "Docente", "proyecto": "Asignatura / Motivo",
            "asiste": "Estado", "tecnico": "Técnico", "observaciones": "Observaciones",
        }
        df_exportar = df[list(columnas_reporte)].rename(columns=columnas_reporte)
        df_exportar["Estado"] = df_exportar["Estado"].map(
            {"Si": "Asistió", "No": "No asistió"}
        ).fillna("Pendiente")
        laboratorio_reporte = "Todos los laboratorios" if lab_filter == "Todos" else LABS_NAMES.get(lab_filter, lab_filter)
        indicadores, series, ficha = analizar_registros(df_exportar, "Fecha", "Estado", fecha_desde, fecha_hasta, "Laboratorio", asistencia=True)
        mostrar_analitica(indicadores, series)
        excel_docentes = _crear_excel_institucional(
            df_exportar,
            "Reporte de asistencia docente",
            "Seguimiento institucional de clases en laboratorios",
            [
                ("Periodo", f"{fecha_desde.strftime('%d/%m/%Y')} al {fecha_hasta.strftime('%d/%m/%Y')}"),
                ("Laboratorio", laboratorio_reporte),
                ("Total de clases", len(df_exportar)),
                ("Asistieron", int((df["asiste"] == "Si").sum())),
                ("No asistieron", int((df["asiste"] == "No").sum())),
            ],
        )
        excel_docentes = completar_excel(excel_docentes, indicadores, series, ficha,
                                         filtros=[("Salón", laboratorio_reporte), ("Fuente", "Registros activos con código PROFESOR y banco cero; incluye asistencia registrada de clases fijas.")])
        st.download_button(
            "Descargar Excel institucional",
            excel_docentes,
            f"reporte_docentes_{fecha_desde.strftime('%Y%m%d')}_{fecha_hasta.strftime('%Y%m%d')}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="doc_descargar_reporte",
        )
        
        # Resumen
        st.subheader(" Resumen de asistencias de docentes")
        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric("Total de clases registradas", len(df))
        with col2:
            asistencias = len(df[df['asiste'] == 'Si'])
            st.metric("Asistieron", asistencias)
        with col3:
            no_asistencias = len(df[df['asiste'] == 'No'])
            st.metric("No asistieron", no_asistencias)
