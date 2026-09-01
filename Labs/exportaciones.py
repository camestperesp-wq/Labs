from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def crear_excel_institucional(df, titulo, subtitulo, metadatos=None, nombre_hoja="Reporte"):
    """Construye un libro ejecutivo, filtrable e imprimible para auditoría."""
    metadatos = metadatos or []
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_hoja[:31]
    total_columnas = max(1, len(df.columns))
    ultima_columna = get_column_letter(total_columnas)

    hoja.merge_cells(f"A1:{ultima_columna}1")
    hoja["A1"] = "Universidad Distrital Francisco José de Caldas"
    hoja["A1"].font = Font(name="Aptos", size=15, bold=True, color="FFFFFF")
    hoja["A1"].fill = PatternFill("solid", fgColor="731116")
    hoja["A1"].alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[1].height = 28

    for fila, texto, tamano, color in (
        (2, titulo, 13, "731116"),
        (3, subtitulo, 10, "5F6368"),
    ):
        hoja.merge_cells(f"A{fila}:{ultima_columna}{fila}")
        hoja.cell(fila, 1, texto)
        hoja.cell(fila, 1).font = Font(name="Aptos", size=tamano, bold=fila == 2, color=color)
        hoja.cell(fila, 1).alignment = Alignment(horizontal="center")

    fila = 5
    for etiqueta, valor in metadatos:
        hoja.cell(fila, 1, etiqueta).font = Font(name="Aptos", bold=True, color="731116")
        hoja.cell(fila, 2, str(valor))
        fila += 1
    hoja.cell(fila, 1, "Generado").font = Font(name="Aptos", bold=True, color="731116")
    hoja.cell(fila, 2, datetime.now().strftime("%d/%m/%Y %H:%M"))
    encabezado_fila = fila + 2

    borde = Border(bottom=Side(style="thin", color="D9DDE2"))
    for columna, nombre in enumerate(df.columns, start=1):
        celda = hoja.cell(encabezado_fila, columna, str(nombre))
        celda.font = Font(name="Aptos", bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="9F1D24")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for fila_excel, valores in enumerate(df.itertuples(index=False, name=None), start=encabezado_fila + 1):
        relleno = PatternFill("solid", fgColor="F3F4F6" if fila_excel % 2 == 0 else "FFFFFF")
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila_excel, columna, "" if pd.isna(valor) else valor)
            celda.font = Font(name="Aptos", size=10, color="292326")
            celda.fill = relleno
            celda.border = borde
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    for indice, columna in enumerate(df.columns, start=1):
        valores = [str(columna)] + ["" if pd.isna(v) else str(v) for v in df[columna].head(250)]
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(len(v) for v in valores) + 3, 42)
    hoja.freeze_panes = f"A{encabezado_fila + 1}"
    if len(df):
        hoja.auto_filter.ref = f"A{encabezado_fila}:{ultima_columna}{encabezado_fila + len(df)}"
    hoja.sheet_view.showGridLines = False
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.oddHeader.center.text = "&B" + titulo
    hoja.oddFooter.center.text = "Página &P de &N"
    hoja.print_title_rows = f"{encabezado_fila}:{encabezado_fila}"

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
