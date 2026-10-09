import io
from datetime import date
import unittest

import pandas as pd
from openpyxl import load_workbook

from analitica_reportes import analizar_registros, completar_excel
from exportaciones import crear_excel_institucional
from constants import LABORATORIOS, LABS_NAMES


class AnaliticaReportesTest(unittest.TestCase):
    def test_promedio_incluye_dias_sin_registros_y_excluye_pendientes_de_tasa(self):
        tabla = pd.DataFrame({"Fecha": ["2026-10-09", "09/10/2026", "2026-10-10"],
                              "Estado": ["Asistió", "No asistió", "Pendiente"],
                              "Salón": ["604", "604", "603"]})
        indicadores, series, ficha = analizar_registros(tabla, "Fecha", "Estado", date(2026, 10, 9), date(2026, 10, 12), "Salón", asistencia=True)
        self.assertEqual(indicadores["Promedio diario de registros"], .75)
        self.assertEqual(indicadores["Asistencia confirmada (%)"], 50)
        self.assertEqual(indicadores["Registros pendientes"], 1)
        self.assertEqual(series["Registros por fecha"]["Registros"].tolist(), [2, 1, 0, 0])
        archivo = completar_excel(crear_excel_institucional(tabla, "Reporte", "Prueba"), indicadores, series, ficha, [("Filtro", "Todos")])
        libro = load_workbook(io.BytesIO(archivo))
        self.assertIn("Ficha técnica", libro.sheetnames)
        self.assertIn("Indicadores", libro.sheetnames)
        capacidades = {fila[0]: fila[1] for fila in libro["Capacidades por salón"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(capacidades, {LABS_NAMES[lab]: capacidad for lab, capacidad in LABORATORIOS.items()})
        self.assertEqual(sum(len(hoja._charts) for hoja in libro.worksheets), 3)

    def test_sin_confirmaciones_no_devuelve_cero_de_asistencia(self):
        tabla = pd.DataFrame({"Fecha": ["2026-10-09"], "Estado": ["Pendiente"]})
        indicadores, _, _ = analizar_registros(tabla, "Fecha", "Estado", date(2026, 10, 9), date(2026, 10, 9), asistencia=True)
        self.assertIsNone(indicadores["Asistencia confirmada (%)"])


if __name__ == "__main__":
    unittest.main()
