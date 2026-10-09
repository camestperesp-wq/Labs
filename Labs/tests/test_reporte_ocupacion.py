from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from streamlit.testing.v1 import AppTest
import streamlit as st

import database as db
import horario_fijo as hf
import reservas
import reportes
from constants import TECNICOS


class ReporteOcupacionTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_patch = patch.object(db, "DB_PATH", Path(self.temp.name) / "ocupacion.db")
        self.db_patch.start()
        db.init_db()
        st.session_state.pop(reservas.CLAVE_INTERCAMBIOS_RESERVAS, None)
        hf.set_horario_celda("Viernes", "08:00-10:00", "604", "Circuitos", "Ingeniería", "ANA TORRES", "Docente Uno")
        hf.set_horario_celda("Viernes", "10:00-12:00", "603", "Adicional", "Adicional", "ANA TORRES", "")

    def tearDown(self):
        st.session_state.pop(reservas.CLAVE_INTERCAMBIOS_RESERVAS, None)
        db.clear_cache()
        self.db_patch.stop()
        self.temp.cleanup()

    def tabla(self):
        return reportes.obtener_reporte_ocupacion(date(2026, 10, 9), date(2026, 10, 9))

    def test_incluye_programacion_sin_reservas_y_asistencia_registrada(self):
        tabla = self.tabla()
        clase = tabla[(tabla["Salón"] == "ELEB 604") & (tabla["Hora"] == "08:00-10:00")].iloc[0]
        self.assertEqual(clase["Tipo"], "Clase fija")
        self.assertEqual(clase["Asistencia docente"], "Pendiente")
        self.assertEqual(clase["Bancos ocupados o bloqueados"], 8)
        adicional = tabla[(tabla["Salón"] == "ELEA 603") & (tabla["Hora"] == "10:00-12:00")].iloc[0]
        self.assertEqual(adicional["Tipo"], "Adicional")
        self.assertEqual(adicional["Bancos disponibles"], 9)
        self.assertEqual(db.ejecutar("SELECT COUNT(*) FROM reservas", fetch=True)[0][0], 0)
        reservas.registrar_asistencia_docente("2026-10-09", "08:00-10:00", "604", "Docente Uno", "Circuitos", "No", TECNICOS[0])
        tabla = self.tabla()
        clase = tabla[(tabla["Salón"] == "ELEB 604") & (tabla["Hora"] == "08:00-10:00")].iloc[0]
        self.assertEqual(clase["Asistencia docente"], "No asistió")
        self.assertEqual(clase["Bancos ocupados o bloqueados"], 0)

    def test_intercambio_respeta_capacidad_fisica_y_monitor_por_fecha(self):
        reservas.asignar_monitor_sesion("2026-10-09", "08:00-10:00", "604", "ANA TORRES")
        reservas.intercambiar_espacios_sesion("2026-10-09", "08:00-10:00", "604", "08:00-10:00", "603")
        tabla = self.tabla()
        clase = tabla[(tabla["Salón"] == "ELEA 603") & (tabla["Hora"] == "08:00-10:00")].iloc[0]
        self.assertEqual(clase["Tipo"], "Clase fija")
        self.assertEqual(clase["Capacidad física"], 9)
        self.assertEqual(clase["Bancos disponibles"], 1)
        self.assertEqual(clase["Monitor"], "ANA TORRES")

    def test_interfaz_genera_tabla_y_excel(self):
        app = AppTest.from_string("from reportes import mostrar_reporte_ocupacion\nmostrar_reporte_ocupacion()").run(timeout=20)
        app.date_input[0].set_value(date(2026, 10, 9))
        app.date_input[1].set_value(date(2026, 10, 9))
        app.button[0].click().run(timeout=20)
        self.assertFalse(app.exception)
        self.assertEqual(len(app.dataframe), 1)
        self.assertEqual(len(app.dataframe[0].value), 2)
        self.assertEqual(len(app.get("download_button")), 1)
        self.assertEqual(len(app.metric), 4)
        self.assertEqual(len(app.get("vega_lite_chart")), 3)

    def test_indices_ponderados_pendientes_y_excel_con_graficas(self):
        import io
        import pandas as pd
        from openpyxl import load_workbook
        from exportaciones import crear_excel_institucional
        tabla = pd.DataFrame({
            "Fecha": ["2026-10-09"] * 3,
            "Hora": ["08:00-10:00"] * 3,
            "Salón": ["A", "B", "C"],
            "Capacidad física": [8, 12, 10],
            "Bancos ocupados o bloqueados": [8, 0, 0],
            "Asistencia docente": ["Asistió", "No asistió", "Pendiente"],
        })
        resumen, series = reportes.resumir_ocupacion(tabla)
        self.assertAlmostEqual(resumen["Índice de ocupación (%)"], 100 * 8 / 30)
        self.assertAlmostEqual(resumen["Promedio de bancos por bloque"], 8 / 3)
        self.assertAlmostEqual(resumen["Bloques con ocupación (%)"], 100 / 3)
        self.assertEqual(resumen["Asistencia docente (%)"], 50)
        self.assertEqual(resumen["Asistencias docentes pendientes"], 1)
        base = crear_excel_institucional(tabla, "Ocupación", "Prueba")
        archivo = reportes._agregar_estadisticas_excel(base, resumen, series)
        libro = load_workbook(io.BytesIO(archivo))
        self.assertIn("Indicadores", libro.sheetnames)
        for nombre in ("Por salón", "Por fecha", "Por hora"):
            self.assertEqual(len(libro[nombre]._charts), 1)

    def test_reservas_y_docentes_generan_analitica(self):
        reservas.registrar_asistencia_docente("2026-10-09", "08:00-10:00", "604", "Docente Uno", "Circuitos", "Si", TECNICOS[0])
        for funcion in ("mostrar_reporte_completo", "mostrar_reporte_asistencia_docentes"):
            app = AppTest.from_string(f"from reportes import {funcion}\n{funcion}()").run(timeout=20)
            app.date_input[0].set_value(date(2026, 10, 9))
            app.date_input[1].set_value(date(2026, 10, 9))
            app.button[0].click().run(timeout=20)
            self.assertFalse(app.exception)
            self.assertEqual(len(app.get("vega_lite_chart")), 3)
            self.assertEqual(len(app.get("download_button")), 1)


if __name__ == "__main__":
    unittest.main()
